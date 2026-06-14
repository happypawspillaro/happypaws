# -*- coding: utf-8 -*-
"""Llena el formulario FR11 (Informe del Director del Proyecto de Titulación)
con los datos del proyecto Happy Paws Píllaro, respetando el formato del
documento (Cambria, marcas 'X' centradas en las casillas SÍ)."""

import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ARCHIVO = "11. FR Informe del proyecto_VA-DAC-0.1.0-FR11.docx"

CARRERA = "Tecnología Superior en Desarrollo de Software"
TEMA = ("Desarrollo e implementación de un sistema web para la gestión de "
        "adopción responsable, casos médicos con recaudación y reportes "
        "comunitarios de la fundación Happy Paws Píllaro, del cantón Santiago "
        "de Píllaro")
AUTOR = "Nicolas Bohórquez Escobar"
CEDULA = "1726220484"
TITULO = "Tecnólogo Superior en Desarrollo de Software"
TUTOR = "Ing. Irving López"
CARGO = "Director del Proyecto de Titulación"
CIUDAD_FECHA = "Quito, 2 de junio del 2026"
FECHA = "2 de junio del 2026"


def set_cell(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, font="Cambria", size=11):
    """Reemplaza el contenido de una celda preservando un formato consistente."""
    # eliminar parrafos extra del template, conservando el primero
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)
    p = cell.paragraphs[0]
    # limpiar runs existentes
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    p.alignment = align
    lines = text.split("\n")
    for i, line in enumerate(lines):
        if i > 0:
            p.add_run().add_break()
        run = p.add_run(line)
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold


def main():
    d = docx.Document(ARCHIVO)

    # Encabezado: fecha de la carta
    for p in d.paragraphs:
        if p.text.strip().startswith("Ciudad,"):
            for r in list(p.runs):
                r._element.getparent().remove(r._element)
            run = p.add_run(CIUDAD_FECHA)
            run.font.name = "Calibri"
            run.font.size = Pt(11)
            break

    # Tabla 0: datos del proyecto
    t0 = d.tables[0]
    set_cell(t0.rows[0].cells[1], CARRERA)
    set_cell(t0.rows[1].cells[1], TEMA)
    set_cell(t0.rows[2].cells[1], AUTOR)
    set_cell(t0.rows[2].cells[3], CEDULA)
    set_cell(t0.rows[3].cells[1], TITULO)

    # Tabla 1: marcar SÍ (columna 1) en las 3 filas de detalle
    t1 = d.tables[1]
    for ri in (1, 2, 3):
        set_cell(t1.rows[ri].cells[1], "X", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Tabla 2: marcar SÍ (columna 1) en las filas de parámetros (2..6)
    t2 = d.tables[2]
    for ri in range(2, 7):
        set_cell(t2.rows[ri].cells[1], "X", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Tabla 3: bloque ELABORADO (firma del director)
    t3 = d.tables[3]
    set_cell(t3.rows[1].cells[0], f"{TUTOR}\n{CARGO}", align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(t3.rows[2].cells[0], FECHA, align=WD_ALIGN_PARAGRAPH.CENTER)

    d.save(ARCHIVO)
    print("FR11 llenado ->", ARCHIVO)


if __name__ == "__main__":
    main()
